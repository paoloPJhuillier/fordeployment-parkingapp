"""
Generates /app/docs/output/Defect_Validation_Response_2026-04-29.xlsx —
v2 of the line-by-line response. Reflects the deeper fixes shipped after
the customer's follow-up screenshots:

  • PPA-54 — date picker — replaced react-day-picker range mode (which
              dropped clicks-inside-existing-range) with two separate
              single-date pickers; every click commits.
  • PPA-62 — admin cancel endpoint also now blocks past-dated cancellations
              (the user-side fix had missed the admin path); UI Cancel
              button hidden on past rows in AdminReservations.
  • PPA-66 — Join Waitlist UX *actually* working end-to-end. Required
              three fixes: parking-config now exposes waitlist_enabled to
              non-admins (V-02 boundary preserved); Couchbase abstraction
              find_one() now accepts sort kwarg (Mongo-API parity);
              event_blocks now factor into is_available so all-fully-booked
              detection fires correctly.
  • NEW    — Block Slots dialog now flags slots already taken by another
              event block (amber "Blocked" tile + reason tooltip).
"""
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path("/app/docs/output/Defect_Validation_Response_2026-04-29.xlsx")
OUT.parent.mkdir(parents=True, exist_ok=True)

STATUS_COLORS = {
    "Resolved":          "27AE60",
    "Already Resolved":  "16A085",
    "Verify Deploy":     "1ABC9C",
    "Verify (UI)":       "2980B9",
    "Backlog (P1)":      "E67E22",
    "Cannot Reproduce":  "7F8C8D",
    "Open":              "C0392B",
}
SEV_COLORS = {"High": "C0392B", "Medium": "E67E22", "Low": "F1C40F"}

# (Defect ID, Title, Module, Severity, Customer Status, Resolution Status,
#  Remarks / Resolution detail, Code Path / Files, QA Verification Steps)
ITEMS = [
    ("PPA-1", "Login - Invalid credentials shows blank page", "Authentication", "High", "Failed",
     "Already Resolved",
     "Closed in earlier remediation cycle. LoginPage.js does a client-side email regex pre-check, the form has noValidate, and the catch block ALWAYS shows a sonner toast and stays on the page.",
     "/app/frontend/src/pages/LoginPage.js (handleLogin)",
     "Try login with malformed email OR unknown email + valid password — expect toast + form remains visible."),

    ("PPA-14", "Vehicle Management - cannot prevent delete with active booking", "Vehicle Management", "High", "Failed",
     "Resolved",
     "DELETE /api/vehicles/{id} now refuses when the vehicle has any reservation in {pending, confirmed} with date >= today, returning 400 'Cannot delete vehicle — it has an active reservation on YYYY-MM-DD'.",
     "/app/backend/routes/vehicles.py (delete_vehicle)",
     "1) Create vehicle V; 2) Create future reservation using V; 3) DELETE /api/vehicles/V → expect 400; 4) Cancel reservation → DELETE succeeds."),

    ("PPA-15", "Building names not aligned in Reservations by Buildings graph", "Admin Dashboard", "Low", "Failed",
     "Already Resolved",
     "YAxis widened to 180px with interval={0} and a tickFormatter that truncates with ellipsis; full name preserved in the recharts Tooltip.",
     "/app/frontend/src/pages/admin/AdminDashboard.js (BarChart YAxis)",
     "Reload Admin Dashboard with at least 4 buildings of varying name lengths — labels should not overlap or be clipped."),

    ("PPA-17", "User Dashboard - tooltip overlaps page title", "User Dashboard", "Low", "Failed",
     "Resolved",
     "Universal fix on the shadcn TooltipContent component — collisionPadding={16}, avoidCollisions=true. Tooltips now flip side near the page edge / sticky header.",
     "/app/frontend/src/components/ui/tooltip.jsx",
     "Hover any info icon near the top of any page — tooltip should flip downward instead of overlapping the header."),

    ("PPA-19", "Daily Reservations chart - dates jumbled / unsorted", "Reports & Analytics", "Medium", "Failed",
     "Resolved",
     "Backend /api/reports/stats now returns daily_breakdown sorted ascending by date.",
     "/app/backend/routes/reports.py (line 91)",
     "GET /api/reports/stats and check daily_breakdown[].date is monotonically increasing."),

    ("PPA-22", "Building Management - slot color does not change to gray when reservation is created", "Building Management", "High", "Failed/Blocked",
     "Already Resolved",
     "BuildingManagement.js refetches on window focus, document visibilitychange, and every 30 s while the tab is visible. Slot colors update without F5.",
     "/app/frontend/src/pages/admin/BuildingManagement.js",
     "Open Building Management in tab A; create reservation in tab B; switch back — slot color flips from green to gray within 1 s of focus."),

    ("PPA-24", "Disconnect between Reservations by Buildings graph and slot statuses", "Reports & Analytics", "Medium", "Failed",
     "Verify Deploy",
     "Same root cause as PPA-22 (data freshness). Polling/focus refetch from PPA-22 + newest-first reservations sort from PPA-57 mean the graph and slot grid reconcile within 30 s. If observation persists post-deploy, invalidate CDN/edge cache.",
     "/app/frontend/src/pages/admin/Reports.js, /app/frontend/src/pages/admin/BuildingManagement.js",
     "Create a reservation; within 30 s reload Reports + Building Management — both should reflect the new booking."),

    ("PPA-30", "User does not receive email notifications for waitlist / reservation events", "Notifications", "Medium", "Failed",
     "Backlog (P1)",
     "Email notifications is an explicit P1 future task. Backend currently writes in-app notifications + has notify_next_waitlisted_user hooks. Bringing this to GA needs an SMTP/Resend/SendGrid integration. Not in scope for this defect cycle.",
     "Future: /app/backend/services/email.py (TBD)",
     "Out of scope for this validation cycle. Item is tracked in PRD.md backlog."),

    ("PPA-31", "Save button hidden under 'Made with Emergent' badge", "Admin Settings", "Medium", "Failed",
     "Resolved",
     "AdminLayout main content area now has pb-24/pb-28 (96 / 112 px bottom padding). Floating Emergent badge no longer covers Save / submit buttons.",
     "/app/frontend/src/pages/admin/AdminLayout.js (lines 154-157)",
     "Open AdminSettings on a viewport ≥ md and verify the Save button is fully clickable; the badge no longer overlaps."),

    ("PPA-54", "Reports - date picker only allows changing END date", "Reports & Analytics", "Medium", "Failed",
     "Resolved",
     "DEEPER FIX (2026-04-29): The earlier QAT round only handled half-range clicks. Customer screenshot showed a residual bug — clicking INSIDE an already-complete range made react-day-picker emit onSelect(undefined), dropping the click. Replaced the single mode='range' Calendar with TWO separate mode='single' pickers (Start / End). Every click commits immediately. Cross-validation: a Start after current End auto-pushes End forward; End picker disables dates before Start.",
     "/app/frontend/src/pages/admin/Reports.js (lines 286-336)",
     "Open Reports → click Start: pick any date → button label updates immediately. Click End: pick any date ≥ Start → updates. Stats refetch on each change."),

    ("PPA-57", "Recent Reservations table shows stale data", "Reports & Analytics", "Medium", "Failed",
     "Already Resolved",
     "Reports.js sorts client-side by created_at DESC and refetches on focus / visibility change. Recent Reservations always shows the latest activity.",
     "/app/frontend/src/pages/admin/Reports.js (fetchReservations + useEffect)",
     "Create a reservation; switch to Reports tab → it should appear at the top of Recent Reservations within 1 s of focus."),

    ("PPA-58", "User Management - tooltip overlaps with header", "User Management", "Low", "Failed",
     "Resolved",
     "Same universal fix as PPA-17 — shadcn TooltipContent uses Radix collisionPadding=16 and avoidCollisions=true.",
     "/app/frontend/src/components/ui/tooltip.jsx",
     "Hover the User Management header tooltip — it should flip downward instead of clashing with the page title."),

    ("PPA-59", "Block Slots dialog - only first conflicting slot reported", "Event Blocking", "Medium", "Failed",
     "Already Resolved",
     "create_event_block does a two-pass create: collect ALL conflicting slots in pass 1, raise once with the full comma-separated list, only insert in pass 2 if pass 1 was empty.",
     "/app/backend/routes/event_blocks.py (create_event_block)",
     "POST /api/event-blocks with 3 slot_ids that all conflict on the same date — error detail must list all 3 slot labels."),

    ("PPA-62", "Booking - able to cancel past reservations", "Reservations", "Medium", "Failed",
     "Resolved",
     "DEEPER FIX (2026-04-29): Earlier round only fixed user-side endpoint. Customer screenshot showed Admin Reservations UI still rendered Cancel button on past rows AND the admin cancel endpoint had no guard. Added the same past-date guard to PUT /api/admin/reservations/{id}/cancel (returns 400 'Cannot cancel a past reservation') AND the AdminReservations table now hides Cancel for any row with date < today. Both endpoints + UI verified.",
     "/app/backend/routes/users.py (admin_cancel_reservation), /app/backend/routes/reservations.py (cancel_reservation), /app/frontend/src/pages/admin/AdminReservations.js (canCancel)",
     "Pick a past reservation in {pending, confirmed}: 1) PUT /api/admin/reservations/{id}/cancel → 400; 2) PUT /api/reservations/{id}/cancel → 400; 3) AdminReservations table — Cancel button NOT shown."),

    ("PPA-64", "Building Management - slot label / status overlap", "Building Management", "Low", "Failed",
     "Verify (UI)",
     "Slot label rendering uses truncate + max-w-full and a title attribute — overflow gets ellipsis, full label on hover. If overlap persists at a specific viewport, share screen size.",
     "/app/frontend/src/pages/admin/BuildingManagement.js (lines 500-503)",
     "Test on viewports ≥ 1280 px and < 768 px. Report any specific overlap with a screenshot."),

    ("PPA-65", "Building Management - 'Active' character-limit indicator behavior", "Building Management", "Low", "Failed",
     "Already Resolved",
     "Phase 6 added explicit maxLength on building text inputs; PUT now uses a permissive BuildingUpdate so legacy buildings round-trip without 422.",
     "/app/backend/models/building.py, /app/frontend/src/pages/admin/BuildingManagement.js",
     "Try editing a legacy building whose name exceeds 30 chars — PUT should succeed; new buildings still constrained."),

    ("PPA-66", "Waitlist - Join Waitlist UX missing from booking flow", "Waitlist", "Medium", "Failed",
     "Resolved",
     "DEEPER FIX (2026-04-29): Customer reported waitlist still didn't surface even with waitlist_enabled=true. Three real bugs were blocking the flow:\n"
     "(1) /api/parking-config public response was hiding waitlist_enabled (V-02 fix had overcorrected) → BookingPage saw config.waitlist_enabled=undefined and never rendered the banner. Fixed: ParkingConfigPublicResponse now exposes UX flags (waitlist_enabled, no_show_release_*, main_building_exclusive); admin-only operational tuning kept hidden.\n"
     "(2) /api/slots/available ignored event_blocks → all 8 Floor-2 slots looked 'available' so allSlotsFullyBooked never fired. Fixed: get_available_slots merges event_blocks into the hourly timeline, returns is_event_blocked + reason.\n"
     "(3) Couchbase abstraction's find_one() didn't accept sort kwarg, but routes/waitlist.py uses it to compute the next queue position → POST /waitlist/join returned 500 on Couchbase backend. Fixed: find_one(filter, projection, sort) now mirrors Motor's signature.",
     "/app/backend/models/parking_config.py, /app/backend/routes/parking_config.py, /app/backend/routes/buildings.py, /app/backend/database/couchbase_db.py",
     "1) Set parking_configs.waitlist_enabled=true; 2) Block all slots for a date (event block + reservations); 3) Open BookingPage as a non-admin user → 'All slots are fully booked / Join Waitlist' banner appears; 4) Click Join → toast + Position #1 badge; 5) Leave → toast + Join button reappears."),

    ("NEW-EVT-BLOCK", "Block Slots dialog does not show existing event-block status", "Event Blocking", "Medium", "Found in re-test",
     "Resolved",
     "NEW FIX (2026-04-29): Customer screenshot showed admin opening Block Slots dialog while another event block was already in effect — slots already taken by the existing block were rendered as green Available (admins could double-block). Fixed by surfacing is_event_blocked + event_block_reason from /api/slots/available; EventBlocking dialog now renders blocked slots in amber with 'Blocked' sublabel, disabled selection, and reason as tooltip. Added 'Event Blocked' to the legend.",
     "/app/backend/routes/buildings.py (get_available_slots), /app/frontend/src/pages/admin/EventBlocking.js",
     "1) Create event block on slots [9, 10, 11] for a date; 2) Open Block Slots dialog for the same building/floor/date; 3) Slots 9, 10, 11 should render amber, label 'Blocked', tooltip = original reason, disabled."),

    ("PPA-14b", "Reset User's Main Building - 'Building not found' error", "User Management", "Medium", "Failed",
     "Cannot Reproduce",
     "Could not reproduce on current build — PUT /api/users/{id} accepts main_building=null and assigned_buildings=[] cleanly. Likely stale frontend state OR the user's current main_building no longer exists. Recommend re-test on freshly deployed build with the exact failing payload captured.",
     "/app/backend/routes/users.py (update_user)",
     "Re-run with browser DevTools network tab open and capture the failing PUT request body + 4xx response."),

    ("PPA-14c", "Block a slot from Building Management - status overlaps with name", "Building Management", "Low", "Failed",
     "Verify (UI)",
     "Same as PPA-64 — slot rendering uses truncate. Need a screenshot at a specific viewport size to reproduce.",
     "/app/frontend/src/pages/admin/BuildingManagement.js",
     "Reproduce with screenshot + viewport width."),

    ("PPA-14d", "View Admin's Dashboard (duplicate row, no detail)", "Admin Dashboard", "—", "Failed",
     "Cannot Reproduce",
     "Row in source spreadsheet had no Title or Actual Result populated — appears to be a duplicate / data-entry artifact. Need QA to specify what was being validated.",
     "—",
     "Please clarify the test scenario; row is currently empty."),

    ("PPA-14e", "View Admin's Dashboard (duplicate row, no detail)", "Admin Dashboard", "—", "Failed",
     "Cannot Reproduce",
     "Same as above — duplicate empty row.",
     "—",
     "Please clarify the test scenario; row is currently empty."),
]


# ---------- Build workbook ----------
wb = Workbook()
ws = wb.active
ws.title = "Defect Response"

HEADERS = ["Defect ID", "Title", "Module", "Severity", "Customer Status",
           "Resolution Status", "Remarks / Resolution Detail",
           "Code Path / Files", "QA Verification Steps"]

hdr_fill = PatternFill("solid", fgColor="08263E")
hdr_font = Font(color="FFFFFF", bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
thin = Side(border_style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = center
    cell.border = border

for r, row in enumerate(ITEMS, 2):
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.alignment = left_wrap if c >= 7 else center
        cell.border = border
    sev = row[3]
    sc = ws.cell(row=r, column=4)
    sc.fill = PatternFill("solid", fgColor=SEV_COLORS.get(sev, "DDDDDD"))
    sc.font = Font(color="FFFFFF", bold=True)
    sc.alignment = center
    rstat = row[5]
    rc = ws.cell(row=r, column=6)
    rc.fill = PatternFill("solid", fgColor=STATUS_COLORS.get(rstat, "DDDDDD"))
    rc.font = Font(color="FFFFFF", bold=True)
    rc.alignment = center

WIDTHS = [16, 50, 22, 11, 18, 18, 80, 55, 55]
for i, w in enumerate(WIDTHS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 32
for r in range(2, len(ITEMS) + 2):
    # Bigger rows for the entries that now carry deeper multi-paragraph remarks.
    item = ITEMS[r - 2]
    ws.row_dimensions[r].height = 220 if item[0] in ("PPA-66", "PPA-62", "PPA-54", "NEW-EVT-BLOCK") else 130
ws.freeze_panes = "A2"

# ---- Summary tab ----
sm = wb.create_sheet("Summary")
sm["A1"] = "Defect Validation — Response Summary (v2)"
sm["A1"].font = Font(size=16, bold=True, color="08263E")
sm["A2"] = (
    f"Source: 'Parking App - Defect Validation.xlsx'   |   "
    f"Generated {datetime.now().strftime('%B %d, %Y')}"
)
sm["A2"].font = Font(italic=True, color="666666")
sm["A3"] = (
    "v2 supersedes the 2026-04-28 file. Rows PPA-54, PPA-62, PPA-66 were "
    "upgraded after the customer's screenshot follow-ups exposed deeper bugs "
    "that the first-pass fixes had missed. NEW-EVT-BLOCK was added during "
    "re-testing."
)
sm["A3"].font = Font(italic=True, color="666666")
sm["A3"].alignment = Alignment(wrap_text=True, vertical="top")

sm["A5"] = "Total customer-reported defect rows + 1 new finding"
sm["B5"] = len(ITEMS)
sm["A5"].font = Font(bold=True)

sm["A7"] = "Resolution status breakdown"
sm["A7"].font = Font(bold=True)
i = 8
for st in ["Resolved", "Already Resolved", "Verify Deploy", "Verify (UI)",
           "Backlog (P1)", "Cannot Reproduce", "Open"]:
    cell = sm.cell(row=i, column=1, value=st)
    cell.fill = PatternFill("solid", fgColor=STATUS_COLORS[st])
    cell.font = Font(color="FFFFFF", bold=True)
    sm.cell(row=i, column=2, value=sum(1 for x in ITEMS if x[5] == st))
    i += 1

i += 1
sm.cell(row=i, column=1,
        value="Items genuinely resolved this cycle (new code)").font = (
    Font(bold=True, color="27AE60"))
i += 1
for x in ITEMS:
    if x[5] == "Resolved":
        sm.cell(row=i, column=1, value=f"{x[0]} — {x[1]}").font = Font(bold=True)
        sm.cell(row=i, column=2, value=x[3])
        i += 1

i += 2
sm.cell(row=i, column=1,
        value="Deeper fixes shipped after customer follow-up (2026-04-29):"
        ).font = Font(bold=True, color="C0392B")
i += 1
for tag, note in [
    ("PPA-54", "Range picker swapped for two single-date pickers (every click commits)."),
    ("PPA-62", "Admin cancel endpoint + UI Cancel button now both block past dates."),
    ("PPA-66", "Three bugs unblocked the Join Waitlist flow (config field hiding + event_block availability + Couchbase find_one sort)."),
    ("NEW-EVT-BLOCK", "Block Slots dialog now flags slots already covered by another event block."),
]:
    sm.cell(row=i, column=1, value=f"{tag} — {note}")
    i += 1

sm.column_dimensions["A"].width = 90
sm.column_dimensions["B"].width = 30

wb.save(OUT)
print(f"Wrote {OUT}")
print(f"Total rows: {len(ITEMS)}")
print(f"  Resolved (this cycle): {sum(1 for x in ITEMS if x[5] == 'Resolved')}")
print(f"  Already resolved (prior): {sum(1 for x in ITEMS if x[5] == 'Already Resolved')}")
print(f"  Verify deploy/UI: {sum(1 for x in ITEMS if x[5] in ('Verify Deploy', 'Verify (UI)'))}")
print(f"  Backlog: {sum(1 for x in ITEMS if x[5] == 'Backlog (P1)')}")
print(f"  Cannot reproduce: {sum(1 for x in ITEMS if x[5] == 'Cannot Reproduce')}")
