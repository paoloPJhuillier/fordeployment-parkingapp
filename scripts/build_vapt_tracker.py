"""
One-shot script to generate /app/docs/output/VAPT_Tracker_2026-03-25.xlsx
from the structured findings list. Re-runnable.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path("/app/docs/output/VAPT_Tracker_2026-03-25.xlsx")
OUT.parent.mkdir(parents=True, exist_ok=True)

# Severity / status colors
SEV_COLORS = {
    "High":   "C0392B",
    "Medium": "E67E22",
    "Low":    "F1C40F",
}
STATUS_COLORS = {
    "Remediated":      "27AE60",
    "Code Fixed - Verify Deploy": "16A085",
    "Not Remediated":  "C0392B",
    "Infra Task":      "7F8C8D",
    "Partial":         "E67E22",
}

# (ID, Title, Severity, VAPT-Status, Real-Status, Affected, Description, Impact, RootCause, RemediationPlan, Owner, Effort, Files)
FINDINGS = [
    (
        "V-01", "Privilege Escalation via Bulk Upload Role Field", "High",
        "Remediated", "Remediated",
        "POST /api/users/bulk-upload",
        "CSV bulk upload accepted role=admin from the file. Caller could create rogue admin accounts at scale.",
        "Compromised admin or insider creates backdoor admins buried in legitimate user batches; no audit trail.",
        "Endpoint trusted client-supplied role.",
        "DONE — backend forces role to 'user' or 'attendant' regardless of CSV; admin role is silently downgraded. Verified end-to-end in retest (22-row CSV with role=admin → all created as 'user').",
        "Backend dev",
        "0.5 day (done)",
        "/app/backend/routes/users.py (bulk_upload + sanitize logic)",
    ),
    (
        "V-02", "Excessive Read Access — /api/parking-config/{id} and /api/slot-registrations", "High",
        "Not Remediated", "Remediated",
        "GET /api/parking-config/{building_id}, GET /api/slot-registrations, GET /api/slot-registrations/user/{id}",
        "Normal users received admin-level fields from parking-config (no_show_release_*, waitlist_*, main_building_exclusive). slot-registrations leaked other users' user_ids in the same building.",
        "Information disclosure; violates least-privilege; aided lateral attacks if a single user account got phished.",
        "Endpoints used Depends(get_current_user) only — no role-based field filtering or row scoping.",
        "DONE Apr 24 2026 — (1) Added ParkingConfigPublicResponse trimmed model; admins still get the full response, non-admins get only release_time, default_start/end_time, booking_window_days. "
        "(2) /api/slot-registrations: users see only their own; attendants see their assigned buildings only. "
        "(3) /api/slot-registrations/user/{user_id}: returns 403 if non-admin requests another user's regs. "
        "Verified via curl: admin gets 11 keys, user gets 6 keys, no leakage of no_show_release_*.",
        "Backend dev",
        "DONE",
        "/app/backend/routes/parking_config.py, /app/backend/routes/building_policies.py, /app/backend/models/parking_config.py",
    ),
    (
        "V-03", "Permissive CORS — Access-Control-Allow-Origin: *", "High",
        "Not Remediated", "Remediated",
        "All /api/* endpoints behind FastAPI CORS middleware",
        "Cross-origin sites could issue authenticated requests if cookies leaked; raised CSRF blast radius.",
        "FastAPI CORS middleware fell back to allow_origins=['*'] when CORS_ORIGINS env var was empty.",
        "DONE Apr 24 2026 — Fallback now [] (fail-closed). Operator must set CORS_ORIGINS env var explicitly for cross-origin deployments. Verified at app layer: direct curl to FastAPI rejects foreign origin (no Access-Control-Allow-Origin header in response). "
        "NOTE: Preview URL still shows '*' because the Emergent K8s ingress proxy adds its own CORS headers — that is a deployment-layer concern. "
        "Recommended for prod: use the new docker-compose nginx same-origin proxy (no CORS at all) — guide in /app/docs/DEPLOYMENT.md.",
        "Backend dev / DevOps",
        "DONE (code); ops follow-up to set CORS_ORIGINS or use proxy",
        "/app/backend/server.py (lines 96-107)",
    ),
    (
        "V-04", "Unminified Source Code Exposed", "Medium",
        "Remediated", "Remediated",
        "Built JS bundles served to the browser",
        "Full API map / role logic visible in DevTools without auth.",
        "Source maps were emitted in production build.",
        "DONE — GENERATE_SOURCEMAP=false in /app/frontend/.env. Verified post-build: App.jsx and api.js no longer readable in DevTools Sources.",
        "Frontend dev",
        "0.25 day (done)",
        "/app/frontend/.env",
    ),
    (
        "V-05", "Decommissioned Environment Still Public — slot-booking-3.emergent.host", "Medium",
        "Not Remediated", "Infra Task",
        "slot-booking-3.emergent.host (old preview URL)",
        "Old build leaks pre-fix source / endpoints; widens attack surface.",
        "Old preview environment was never torn down on the Emergent platform.",
        "Cannot be done from app code. Action items: "
        "(1) From the Emergent dashboard: delete or disable the old preview deployment. "
        "(2) If kept for rollback, move it behind basic-auth or VPN. "
        "(3) Once decommissioned, request rescan to confirm DNS no longer resolves.",
        "Platform owner / DevOps",
        "30 min ops",
        "Emergent platform UI (no code change)",
    ),
    (
        "V-06", "Attendant Role Saw All Buildings", "Medium",
        "Remediated", "Remediated",
        "GET /api/buildings",
        "Attendants saw ALL buildings instead of only their assigned one(s).",
        "Endpoint did not filter by current_user.assigned_buildings.",
        "DONE — buildings.py applies role-based filter (admin = all; attendant/user = assigned + main_building). Retest confirmed restriction.",
        "Backend dev",
        "0.5 day (done)",
        "/app/backend/routes/buildings.py (get_buildings)",
    ),
    (
        "V-07", "API Version Disclosed at /api/", "Low",
        "Remediated", "Remediated",
        "GET /api/",
        "Discloses 'Company Parking API' name + exact version to unauthenticated callers.",
        "Endpoint hand-rolled to return version metadata.",
        "DONE — /api/ now returns {message, status: 'running'}. No version, no internal name. Confirmed in code (server.py line ~88).",
        "Backend dev",
        "10 min (done)",
        "/app/backend/server.py (root)",
    ),
    (
        "V-08", "CSV Formula Injection (Defense-in-Depth) via Bulk Upload", "Low",
        "Not Remediated", "Remediated",
        "POST /api/users/bulk-upload, POST /api/users, PUT /api/users/{id}",
        "If a CSV/Excel export is added later, =CMD()/+CMD()/=HYPERLINK() payloads stored today could RCE on the analyst's machine.",
        "Auditor saw raw payloads stored in the deployed env.",
        "DONE Apr 24 2026 — sanitize_csv_field() now applied on ALL three write paths (single-user POST, single-user PUT, bulk upload) for first_name, last_name, company, job_family, parking_sticker_number. "
        "Backfill script /app/backend/scripts/backfill_csv_sanitization.py written; ran in dry-run mode → 0 tainted rows in current DB (clean). "
        "Verified via curl: POSTing first_name='=CMD()' is stored as \"'=CMD()\" (Excel treats as text).",
        "Backend dev",
        "DONE",
        "/app/backend/routes/users.py (create_user, update_user), /app/backend/scripts/backfill_csv_sanitization.py",
    ),
    (
        "V-09", "Login Email Placeholder Reveals Internal Domain", "Low",
        "Remediated", "Remediated",
        "Login page email input",
        "Discloses '*.cebuana.com' email convention to unauthenticated visitors.",
        "Placeholder text was a sample address.",
        "DONE — placeholder changed to 'Enter your email address' in /app/frontend/src/pages/LoginPage.js.",
        "Frontend dev",
        "5 min (done)",
        "/app/frontend/src/pages/LoginPage.js",
    ),
    (
        "V-10", "Outdated react-router 7.11.0 — multiple CVEs", "Low",
        "Not Remediated", "Remediated",
        "Frontend dependency react-router-dom",
        "Open redirect / XSS / CSRF CVEs exist for 7.11.0 — not exploitable today (Declarative Mode) but risky as routing evolves.",
        "package.json was on 7.11.0. Retest saw 7.13.1 in Chrome but 7.11.0 in Firefox — mixed cache state suggested stale build still served.",
        "DONE Apr 24 2026 — package.json pinned to 7.13.1 (since Phase 7). Added unauthenticated GET /api/system/build-info endpoint that returns deploy timestamp and security-sensitive package versions, so retesters can confirm cache freshness with one curl. "
        "Production cutover MUST: (a) build with --no-cache, (b) invalidate any CDN/edge caches, (c) hard-refresh browsers OR set Cache-Control: no-store on index.html (the new docker-compose nginx config already does this).",
        "Frontend dev / DevOps",
        "DONE (code); ops to invalidate CDN cache on next deploy",
        "/app/frontend/package.json, /app/backend/routes/system.py (build-info endpoint)",
    ),
    (
        "V-11", "Cleartext Password Accessible via DOM/JS Console", "Low",
        "Not Remediated", "Partial",
        "Login page password <input type=\"password\">",
        "While typing, document.getElementById('login-password').value returns the cleartext password. Exploit needs local access, malicious extension, or XSS.",
        "Standard HTML behavior. We use an uncontrolled useRef input and CLEAR the value on submit success/failure, but the ref is still readable while the user is typing.",
        "PARTIAL Apr 24 2026 — Added aggressive wipe-on-blur and visibilitychange handlers in /app/frontend/src/pages/LoginPage.js so the cleartext is purged from the DOM the moment the user tabs away or switches windows. Limits exposure window. "
        "Compensating controls (unchanged): HttpOnly auth cookies, strict CSP, no source maps in prod, X-Frame-Options DENY. "
        "Residual risk accepted: an attacker with active JS injection in the same tab while the user is actively typing can still read .value — this is fundamental to <input type='password'> across the web. Major banks and IDPs accept the same residual risk.",
        "Frontend dev / Security PM",
        "DONE (mitigation); residual accepted",
        "/app/frontend/src/pages/LoginPage.js (wipe useEffect)",
    ),
]

# ----- Build workbook ------------------------------------------------------
wb = Workbook()

# Tab 1 — Tracker
ws = wb.active
ws.title = "Tracker"

HEADERS = [
    "ID", "Title", "Severity",
    "VAPT Retest Status (Mar 25)",
    "Actual Status (Apr 24)",
    "Affected URL/Component",
    "Description",
    "Impact",
    "Root Cause / Why Still Open",
    "Remediation Plan",
    "Owner",
    "Effort",
    "Files / Location",
]

# Header style
hdr_fill = PatternFill("solid", fgColor="08263E")
hdr_font = Font(color="FFFFFF", bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
thin = Side(border_style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c, h in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = center
    cell.border = border

# Rows
for r, row in enumerate(FINDINGS, start=2):
    for c, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.alignment = left_wrap if c >= 6 else center
        cell.border = border
    # Severity color
    sev = row[2]
    sev_cell = ws.cell(row=r, column=3)
    sev_cell.fill = PatternFill("solid", fgColor=SEV_COLORS.get(sev, "DDDDDD"))
    sev_cell.font = Font(color="FFFFFF", bold=True)
    sev_cell.alignment = center
    # VAPT-side status
    vstat = row[3]
    vc = ws.cell(row=r, column=4)
    vc.fill = PatternFill("solid", fgColor=STATUS_COLORS.get(vstat, "DDDDDD"))
    vc.font = Font(color="FFFFFF", bold=True)
    vc.alignment = center
    # Real-side status
    rstat = row[4]
    rc = ws.cell(row=r, column=5)
    rc.fill = PatternFill("solid", fgColor=STATUS_COLORS.get(rstat, "DDDDDD"))
    rc.font = Font(color="FFFFFF", bold=True)
    rc.alignment = center

# Column widths
WIDTHS = [7, 38, 10, 16, 18, 30, 50, 40, 45, 70, 18, 18, 40]
for i, w in enumerate(WIDTHS, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# Row heights
ws.row_dimensions[1].height = 35
for r in range(2, len(FINDINGS) + 2):
    ws.row_dimensions[r].height = 130

# Freeze header
ws.freeze_panes = "A2"

# Tab 2 — Summary
sm = wb.create_sheet("Summary")
sm["A1"] = "VAPT Remediation — Cebuana Parking App"
sm["A1"].font = Font(size=16, bold=True, color="08263E")
sm["A2"] = "Source: 2026-03-25 VAPT report retest"
sm["A2"].font = Font(italic=True, color="666666")
sm["A4"] = "Total findings"
sm["B4"] = len(FINDINGS)
sm["A5"] = "By severity"
for i, sev in enumerate(["High", "Medium", "Low"]):
    sm.cell(row=6 + i, column=1, value=sev).fill = PatternFill("solid", fgColor=SEV_COLORS[sev])
    sm.cell(row=6 + i, column=1).font = Font(color="FFFFFF", bold=True)
    sm.cell(row=6 + i, column=2, value=sum(1 for f in FINDINGS if f[2] == sev))

sm["A11"] = "By actual status (Apr 24)"
sm["A11"].font = Font(bold=True)
statuses = ["Remediated", "Code Fixed - Verify Deploy", "Partial", "Not Remediated", "Infra Task"]
for i, st in enumerate(statuses):
    cell = sm.cell(row=12 + i, column=1, value=st)
    cell.fill = PatternFill("solid", fgColor=STATUS_COLORS.get(st, "DDDDDD"))
    cell.font = Font(color="FFFFFF", bold=True)
    sm.cell(row=12 + i, column=2, value=sum(1 for f in FINDINGS if f[4] == st))

sm["A20"] = "Open items requiring action"
sm["A20"].font = Font(bold=True, color="C0392B")
open_items = [f for f in FINDINGS if f[4] in ("Not Remediated", "Code Fixed - Verify Deploy", "Partial", "Infra Task")]
for i, f in enumerate(open_items):
    sm.cell(row=21 + i, column=1, value=f"{f[0]} — {f[1]}").font = Font(bold=True)
    sm.cell(row=21 + i, column=2, value=f"({f[2]}) {f[4]}")

sm.column_dimensions["A"].width = 70
sm.column_dimensions["B"].width = 35

wb.save(OUT)
print(f"Wrote {OUT}")
print(f"Findings: {len(FINDINGS)}")
print(f"Open: {len(open_items)} of {len(FINDINGS)}")
