#!/usr/bin/env python3
"""
Convert markdown documentation to Word (.docx) and PDF files.
Modern business style with table of contents.
"""

import os
import re
import markdown
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from lxml import etree


# Document configuration
DOCS_DIR = Path("/app/docs")
OUTPUT_DIR = Path("/app/docs/output")
OUTPUT_DIR.mkdir(exist_ok=True)

# Color scheme - Modern Business (Dark Navy + Teal accent)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
TEAL = RGBColor(0x00, 0x96, 0x88)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MED_GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_HEADER_BG = "1B2A4A"
TABLE_ALT_BG = "F0F4F8"

DOCUMENTS = [
    ("01-business-requirements/BRD.md", "01_Business_Requirements"),
    ("02-system-requirements/SRS.md", "02_System_Requirements_Specification"),
    ("03-system-design/system-design.md", "03_System_Design"),
    ("03-system-design/diagrams/architecture.md", "03a_Architecture_Diagrams"),
    ("03-system-design/diagrams/erd.md", "03b_Entity_Relationship_Diagram"),
    ("03-system-design/diagrams/sequences.md", "03c_Sequence_Diagrams"),
    ("03-system-design/compliance-matrix.md", "03d_Compliance_Matrix"),
    ("04-testing/test-plan.md", "04_Test_Plan"),
    ("04-testing/test-scripts.md", "04a_Test_Scripts"),
    ("04-testing/test-results.md", "04b_Test_Execution_Results"),
    ("05-user-guides/end-user-guide.md", "05a_End_User_Guide"),
    ("05-user-guides/admin-user-guide.md", "05b_Admin_User_Guide"),
    ("05-user-guides/attendant-guide.md", "05c_Attendant_Guide"),
    ("06-deployment/deployment-guide.md", "06_Deployment_Guide"),
]


def setup_styles(doc):
    """Configure document styles for modern business look."""
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = DARK_GRAY
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15

    for level in range(1, 5):
        style_name = f'Heading {level}'
        if style_name in doc.styles:
            h = doc.styles[style_name]
            h.font.name = 'Calibri'
            h.font.bold = True
            h.font.color.rgb = NAVY
            if level == 1:
                h.font.size = Pt(22)
                h.paragraph_format.space_before = Pt(24)
                h.paragraph_format.space_after = Pt(12)
            elif level == 2:
                h.font.size = Pt(16)
                h.paragraph_format.space_before = Pt(18)
                h.paragraph_format.space_after = Pt(8)
            elif level == 3:
                h.font.size = Pt(13)
                h.font.color.rgb = TEAL
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(6)
            elif level == 4:
                h.font.size = Pt(11)
                h.paragraph_format.space_before = Pt(8)
                h.paragraph_format.space_after = Pt(4)

    # TOC heading style
    if 'TOC Heading' in doc.styles:
        toc_style = doc.styles['TOC Heading']
        toc_style.font.name = 'Calibri'
        toc_style.font.size = Pt(18)
        toc_style.font.color.rgb = NAVY


def add_cover_page(doc, title, subtitle=""):
    """Add a modern cover page."""
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # Spacer
    for _ in range(6):
        doc.add_paragraph("")

    # Accent line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("━" * 40)
    run.font.color.rgb = TEAL
    run.font.size = Pt(14)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = NAVY
    run.font.name = 'Calibri'
    p.paragraph_format.space_after = Pt(4)

    # Subtitle
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(subtitle)
        run.font.size = Pt(14)
        run.font.color.rgb = MED_GRAY
        run.font.name = 'Calibri'

    # Accent line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("━" * 40)
    run.font.color.rgb = TEAL
    run.font.size = Pt(14)

    # Project info
    doc.add_paragraph("")
    info_items = [
        ("Project", "Cebuana Lhuillier Parking Reservation System"),
        ("Version", "1.0"),
        ("Date", "February 22, 2026"),
        ("Classification", "Internal — Confidential"),
    ]
    for label, value in info_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(f"{label}: ")
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = NAVY
        run = p.add_run(value)
        run.font.size = Pt(10)
        run.font.color.rgb = DARK_GRAY
        p.paragraph_format.space_after = Pt(2)

    doc.add_page_break()


def add_toc(doc):
    """Add a Table of Contents field."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Table of Contents")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = NAVY
    run.font.name = 'Calibri'
    p.paragraph_format.space_after = Pt(12)

    # Accent line under TOC heading
    p = doc.add_paragraph()
    run = p.add_run("━" * 30)
    run.font.color.rgb = TEAL
    run.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(12)

    # TOC field code
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_char_begin = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._r.append(fld_char_begin)

    run2 = paragraph.add_run()
    instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>')
    run2._r.append(instr)

    run3 = paragraph.add_run()
    fld_char_sep = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    run3._r.append(fld_char_sep)

    run4 = paragraph.add_run("Right-click and select 'Update Field' to populate the Table of Contents.")
    run4.font.color.rgb = MED_GRAY
    run4.font.italic = True
    run4.font.size = Pt(9)

    run5 = paragraph.add_run()
    fld_char_end = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run5._r.append(fld_char_end)

    doc.add_page_break()


def style_table_cell(cell, bg_color=None, bold=False, font_color=None, font_size=None):
    """Apply styling to a table cell."""
    if bg_color:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.name = 'Calibri'
            if bold:
                run.font.bold = True
            if font_color:
                run.font.color.rgb = font_color
            if font_size:
                run.font.size = font_size


def add_image_to_doc(doc, img_rel_path, alt_text):
    """Add an image to the docx document."""
    # Resolve the image path relative to the current markdown file's directory
    img_path = (CURRENT_MD_DIR / img_rel_path).resolve()
    if not img_path.exists():
        # Try relative to DOCS_DIR
        img_path = (DOCS_DIR / img_rel_path).resolve()
    if not img_path.exists():
        # Fallback: just add the alt text
        p = doc.add_paragraph()
        run = p.add_run(f"[Image: {alt_text}]")
        run.font.italic = True
        run.font.color.rgb = MED_GRAY
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    try:
        run.add_picture(str(img_path), width=Inches(5.5))
    except Exception:
        try:
            run.add_picture(str(img_path), width=Inches(4.0))
        except Exception:
            run = p.add_run(f"[Image: {alt_text}]")
            run.font.italic = True

    # Caption
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(alt_text)
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = MED_GRAY
    cap.paragraph_format.space_after = Pt(8)


# Global variable to track current markdown file's directory
CURRENT_MD_DIR = DOCS_DIR

def parse_markdown_to_docx(md_content, doc):
    """Parse markdown content and add to docx document."""
    lines = md_content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i]

        # Code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                code_text = '\n'.join(code_lines)
                if code_text.strip():
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(6)
                    # Add shading to paragraph
                    pPr = p._p.get_or_add_pPr()
                    shading = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F5F5F5"/>')
                    pPr.append(shading)
                    run = p.add_run(code_text)
                    run.font.name = 'Consolas'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = DARK_GRAY
                code_lines = []
                in_code_block = False
            else:
                # Flush any pending table
                if in_table and table_rows:
                    add_table_to_doc(doc, table_rows)
                    table_rows = []
                    in_table = False
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Image references - ![alt](path)
        img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)', line.strip())
        if img_match:
            if in_table and table_rows:
                add_table_to_doc(doc, table_rows)
                table_rows = []
                in_table = False
            alt_text = img_match.group(1)
            img_path = img_match.group(2)
            add_image_to_doc(doc, img_path, alt_text)
            i += 1
            continue

        # Table rows
        if '|' in line and line.strip().startswith('|'):
            stripped = line.strip()
            # Check if separator row
            if re.match(r'^\|[\s\-:|]+\|$', stripped):
                i += 1
                continue
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            if cells:
                if not in_table:
                    in_table = True
                table_rows.append(cells)
            i += 1
            continue
        else:
            if in_table and table_rows:
                add_table_to_doc(doc, table_rows)
                table_rows = []
                in_table = False

        # Horizontal rule
        if line.strip() in ('---', '***', '___'):
            p = doc.add_paragraph()
            run = p.add_run("━" * 50)
            run.font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
            run.font.size = Pt(6)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,4})\s+(.+)', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            # Clean markdown formatting
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            text = re.sub(r'`(.+?)`', r'\1', text)
            doc.add_heading(text, level=level)
            i += 1
            continue

        # Blockquote
        if line.strip().startswith('>'):
            text = line.strip().lstrip('>').strip()
            text = clean_inline_md(text)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            pPr = p._p.get_or_add_pPr()
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="E8F5E9"/>')
            pPr.append(shading)
            run = p.add_run(text)
            run.font.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = MED_GRAY
            i += 1
            continue

        # Bullet/numbered list
        list_match = re.match(r'^(\s*)([-*+]|\d+\.)\s+(.+)', line)
        if list_match:
            indent = len(list_match.group(1))
            marker = list_match.group(2)
            text = list_match.group(3).strip()
            text = clean_inline_md(text)
            is_numbered = bool(re.match(r'\d+\.', marker))
            style = 'List Number' if is_numbered else 'List Bullet'
            p = doc.add_paragraph(style=style)
            if indent > 0:
                p.paragraph_format.left_indent = Cm(1 + indent * 0.3)
            add_formatted_text(p, text)
            i += 1
            continue

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Navigation links at bottom (skip)
        if line.strip().startswith('*') and ('Previous:' in line or 'Next:' in line or 'Back to:' in line or 'Related:' in line):
            i += 1
            continue

        # Regular paragraph
        text = line.strip()
        if text:
            text = clean_inline_md(text)
            p = doc.add_paragraph()
            add_formatted_text(p, text)

        i += 1

    # Flush remaining table
    if in_table and table_rows:
        add_table_to_doc(doc, table_rows)


def clean_inline_md(text):
    """Remove markdown link syntax, keep text."""
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text


def add_formatted_text(paragraph, text):
    """Add text with bold/italic/code formatting."""
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.font.bold = True
            run.font.name = 'Calibri'
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            run = paragraph.add_run(part[1:-1])
            run.font.italic = True
            run.font.name = 'Calibri'
        elif part:
            run = paragraph.add_run(part)
            run.font.name = 'Calibri'


def add_table_to_doc(doc, rows):
    """Add a styled table to the document."""
    if not rows or not rows[0]:
        return

    num_cols = len(rows[0])
    # Normalize row lengths
    for r in rows:
        while len(r) < num_cols:
            r.append("")

    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'

    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx < num_cols:
                cell = table.rows[row_idx].cells[col_idx]
                cell.text = ""
                p = cell.paragraphs[0]
                text = clean_inline_md(cell_text.strip())
                add_formatted_text(p, text)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)

                if row_idx == 0:
                    # Header row styling
                    style_table_cell(cell, bg_color=TABLE_HEADER_BG, bold=True, font_color=WHITE, font_size=Pt(9.5))
                else:
                    # Alternating row colors
                    bg = TABLE_ALT_BG if row_idx % 2 == 0 else None
                    style_table_cell(cell, bg_color=bg, font_size=Pt(9.5))

    doc.add_paragraph("")  # spacing after table


def add_footer(doc, title):
    """Add footer with page numbers and document title."""
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(f"{title}  |  Cebuana Lhuillier  |  Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = MED_GRAY
    run.font.name = 'Calibri'

    # Page number field
    fld_xml = (
        f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'
    )
    run2 = p.add_run()
    run2._r.append(parse_xml(fld_xml))
    run3 = p.add_run()
    run3._r.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>'))
    run4 = p.add_run()
    run4._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
    run5 = p.add_run("1")
    run5.font.size = Pt(8)
    run5.font.color.rgb = MED_GRAY
    run6 = p.add_run()
    run6._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))


def get_title_from_md(md_content):
    """Extract the first H1 heading as the title."""
    for line in md_content.split('\n'):
        match = re.match(r'^#\s+(.+)', line)
        if match:
            return match.group(1).strip()
    return "Untitled"


def convert_document(md_path, output_name):
    """Convert a single markdown file to docx."""
    md_file = DOCS_DIR / md_path
    if not md_file.exists():
        print(f"  SKIP (not found): {md_path}")
        return None

    md_content = md_file.read_text(encoding='utf-8')
    title = get_title_from_md(md_content)

    # Remove the first H1 line (will be on cover page)
    lines = md_content.split('\n')
    content_lines = []
    found_h1 = False
    for line in lines:
        if not found_h1 and re.match(r'^#\s+', line):
            found_h1 = True
            continue
        content_lines.append(line)
    body_content = '\n'.join(content_lines)

    # Determine subtitle from metadata lines
    subtitle = ""
    for line in content_lines[:5]:
        if line.startswith('**Project:**'):
            subtitle = line.replace('**Project:**', '').strip()
            break

    doc = Document()
    setup_styles(doc)
    add_cover_page(doc, title, subtitle)
    add_toc(doc)

    # Set the current directory context for image resolution
    global CURRENT_MD_DIR
    CURRENT_MD_DIR = md_file.parent

    parse_markdown_to_docx(body_content, doc)
    add_footer(doc, title)

    docx_path = OUTPUT_DIR / f"{output_name}.docx"
    doc.save(str(docx_path))
    print(f"  OK: {docx_path.name} ({docx_path.stat().st_size // 1024}KB)")
    return docx_path


def generate_pdf_from_md(md_path, output_name):
    """Convert markdown to PDF via HTML using WeasyPrint."""
    md_file = DOCS_DIR / md_path
    if not md_file.exists():
        return None

    md_content = md_file.read_text(encoding='utf-8')
    title = get_title_from_md(md_content)

    # Resolve relative image paths to absolute file:// URLs
    md_dir = md_file.parent
    def resolve_img(match):
        alt = match.group(1)
        rel_path = match.group(2)
        abs_path = (md_dir / rel_path).resolve()
        if abs_path.exists():
            return f'![{alt}](file://{abs_path})'
        return match.group(0)
    md_content = re.sub(r'!\[([^\]]*)\]\(([^)]+)\)', resolve_img, md_content)

    # Convert markdown to HTML
    html_body = markdown.markdown(
        md_content,
        extensions=['tables', 'fenced_code', 'toc', 'nl2br'],
        extension_configs={
            'toc': {'title': 'Table of Contents', 'toc_depth': 3}
        }
    )

    # Modern business CSS
    css = """
    @page {
        size: A4;
        margin: 2.5cm;
        @bottom-center {
            content: counter(page);
            font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
            font-size: 9pt;
            color: #999;
        }
    }
    @page :first {
        @bottom-center { content: none; }
    }
    body {
        font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
        font-size: 10.5pt;
        line-height: 1.5;
        color: #333;
    }
    h1 {
        color: #1B2A4A;
        font-size: 22pt;
        border-bottom: 3px solid #009688;
        padding-bottom: 8px;
        margin-top: 30px;
        page-break-after: avoid;
    }
    h2 {
        color: #1B2A4A;
        font-size: 16pt;
        margin-top: 24px;
        border-bottom: 1px solid #E0E0E0;
        padding-bottom: 4px;
        page-break-after: avoid;
    }
    h3 {
        color: #009688;
        font-size: 13pt;
        margin-top: 18px;
        page-break-after: avoid;
    }
    h4 {
        color: #1B2A4A;
        font-size: 11pt;
        margin-top: 12px;
    }
    table {
        width: 100%;
        border-collapse: collapse;
        margin: 12px 0;
        font-size: 9.5pt;
        page-break-inside: auto;
    }
    thead tr {
        background-color: #1B2A4A;
        color: white;
    }
    th {
        padding: 8px 10px;
        text-align: left;
        font-weight: 600;
        border: 1px solid #1B2A4A;
    }
    td {
        padding: 6px 10px;
        border: 1px solid #DDD;
    }
    tr:nth-child(even) {
        background-color: #F0F4F8;
    }
    code {
        font-family: 'Courier New', Consolas, monospace;
        background-color: #F5F5F5;
        padding: 1px 4px;
        border-radius: 3px;
        font-size: 9pt;
        color: #C7254E;
    }
    pre {
        background-color: #F5F5F5;
        padding: 12px 16px;
        border-radius: 4px;
        border-left: 4px solid #009688;
        overflow-x: auto;
        font-size: 8.5pt;
        line-height: 1.4;
        page-break-inside: avoid;
    }
    pre code {
        background: none;
        padding: 0;
        color: #333;
    }
    blockquote {
        border-left: 4px solid #009688;
        margin: 12px 0;
        padding: 8px 16px;
        background-color: #E8F5E9;
        color: #555;
        font-style: italic;
    }
    hr {
        border: none;
        border-top: 1px solid #E0E0E0;
        margin: 20px 0;
    }
    img {
        max-width: 100%;
        display: block;
        margin: 12px auto;
        border: 1px solid #E0E0E0;
        border-radius: 4px;
    }
    .toc {
        background-color: #F8F9FA;
        padding: 16px 24px;
        border-radius: 6px;
        border: 1px solid #E0E0E0;
        margin-bottom: 24px;
    }
    .toc a {
        color: #1B2A4A;
        text-decoration: none;
    }
    .toc a:hover {
        color: #009688;
    }
    .toc ul {
        list-style: none;
        padding-left: 16px;
    }
    .toc > ul {
        padding-left: 0;
    }
    .toc li {
        margin: 4px 0;
        font-size: 10pt;
    }
    strong { color: #1B2A4A; }
    a { color: #009688; text-decoration: none; }
    ul, ol { margin: 8px 0; padding-left: 24px; }
    li { margin: 3px 0; }
    """

    html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{title}</title>
    <style>{css}</style>
</head>
<body>
{html_body}
</body>
</html>"""

    pdf_path = OUTPUT_DIR / f"{output_name}.pdf"
    try:
        from weasyprint import HTML
        HTML(string=html).write_pdf(str(pdf_path))
        print(f"  OK: {pdf_path.name} ({pdf_path.stat().st_size // 1024}KB)")
        return pdf_path
    except Exception as e:
        print(f"  WARN PDF: {output_name} - {e}")
        return None


def main():
    print("=" * 60)
    print("Documentation Compiler")
    print("Cebuana Lhuillier Parking Reservation System")
    print("=" * 60)

    print(f"\nOutput directory: {OUTPUT_DIR}\n")

    print("Generating Word (.docx) documents...")
    print("-" * 40)
    docx_files = []
    for md_path, output_name in DOCUMENTS:
        result = convert_document(md_path, output_name)
        if result:
            docx_files.append(result)

    print(f"\nGenerated {len(docx_files)} Word documents.\n")

    print("Generating PDF documents...")
    print("-" * 40)
    pdf_files = []
    for md_path, output_name in DOCUMENTS:
        result = generate_pdf_from_md(md_path, output_name)
        if result:
            pdf_files.append(result)

    print(f"\nGenerated {len(pdf_files)} PDF documents.\n")

    print("=" * 60)
    print(f"TOTAL: {len(docx_files)} DOCX + {len(pdf_files)} PDF files")
    print(f"Location: {OUTPUT_DIR}")
    print("=" * 60)


if __name__ == "__main__":
    main()
