"""
Generates /app/docs/output/Defect_Validation_Response_2026-04-28.xlsx — a
line-by-line response to the customer's defect validation file. Each row
mirrors the customer's defect ID and adds Resolved? + Remarks + Code Path.
"""
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path("/app/docs/output/Defect_Validation_Response_2026-04-28.xlsx")
OUT.parent.mkdir(parents=True, exist_ok=True)

STATUS_COLORS = {
    "Resolved":          "27AE60",
    "Already Resolved":  "16A085",
    "Verify Deploy":     "1ABC9C",
    "Verify (UI)":       "2980B9",
    "Backlog (P1)":      "E67E22",
    "Cannot Reproduce":  "7F8C8D",
    "Open":              "C0392B",
}
SEV_COLORS = {"High":"C0392B","Medium":"E67E22","Low":"F1C40F"}

# (Defect ID, Title, Module, Severity, Status from QA, ResolutionStatus,
#  Remarks/Resolution detail, Code Path/Files, Verification steps for QA)
ITEMS = [
    ("PPA-1", "Login - Invalid credentials shows blank page", "Authentication", "High", "Failed",
     "Already Resolved",
     "Closed in earlier remediation cycle (TCID-LOGIN-005). LoginPage.js now does a client-side email regex pre-check, the form has `noValidate`, and the catch block ALWAYS shows a sonner toast and stays on the page (never navigates on error).",
     "/app/frontend/src/pages/LoginPage.js (handleLogin)",
     "Try login with malformed email (e.g. 'foo') OR unknown email + valid password — expect toast + form remains visible."),

    ("PPA-14", "Vehicle Management - cannot prevent delete with active booking", "Vehicle Management", "High", "Failed",
     "Resolved",
     "DELETE /api/vehicles/{id} now refuses to delete when the vehicle has any reservation in {pending, confirmed} with date >= today, returning 400 'Cannot delete vehicle — it has an active reservation on YYYY-MM-DD'. Past reservations no longer block deletion.",
     "/app/backend/routes/vehicles.py (delete_vehicle)",
     "1) Create vehicle V; 2) Create future reservation using V; 3) Attempt DELETE /api/vehicles/V → expect 400; 4) Cancel reservation → DELETE succeeds."),

    ("PPA-15", "Building names not aligned in Reservations by Buildings graph", "Admin Dashboard", "Low", "Failed",
     "Already Resolved",
     "Closed in QAT round (TCID-DASHBOARD-008). YAxis widened to 180px with interval={0} and a tickFormatter that truncates with ellipsis; full name preserved in the recharts Tooltip.",
     "/app/frontend/src/pages/admin/AdminDashboard.js (BarChart YAxis)",
     "Reload Admin Dashboard with at least 4 buildings of varying name lengths — labels should not overlap or be clipped."),

    ("PPA-17", "User Dashboard - tooltip overlaps page title", "User Dashboard", "Low", "Failed",
     "Resolved",
     "Universal fix on the shadcn TooltipContent component: now passes collisionPadding={16} and avoidCollisions=true so tooltips automatically flip side when near the page edge / sticky header. Applies to every tooltip in the app.",
     "/app/frontend/src/components/ui/tooltip.jsx",
     "Hover any info icon near the top of any page — tooltip should flip downward instead of overlapping the header."),

    ("PPA-19", "Daily Reservations chart - dates jumbled / unsorted", "Reports & Analytics", "Medium", "Failed",
     "Resolved",
     "Backend /api/reports/stats now returns daily_breakdown sorted ascending by date. Verified via curl — dates come out in chronological order.",
     "/app/backend/routes/reports.py (line 91)",
     "GET /api/reports/stats and check daily_breakdown[].date is monotonically increasing."),

    ("PPA-22", "Building Management - slot color does not change to gray when reservation is created", "Building Management", "High", "Failed/Blocked",
     "Already Resolved",
     "Closed in QAT round (TCID-BUILDING-MANAGEMENT-027). BuildingManagement.js refetches on window focus, on document visibilitychange, and every 30 seconds while the tab is visible. Slot colors update without F5.",
     "/app/frontend/src/pages/admin/BuildingManagement.js",
     "Open Building Management in tab A; create reservation in tab B; switch back to tab A — slot color flips from green to gray within 1 second of focus."),

    ("PPA-24", "Disconnect between Reservations by Buildings graph and slot statuses", "Reports & Analytics", "Medium", "Failed",
     "Verify Deploy",
     "Same root cause as PPA-22 (data freshness). The polling/focus refetch from PPA-22 + the newest-first reservations sort from PPA-57/TCID-ANALYTICS-016 mean the graph and the slot grid now reconcile within 30 seconds. If the QA observation persists post-deploy, the next likely cause is a CDN/edge cache — invalidate and retest.",
     "/app/frontend/src/pages/admin/Reports.js, /app/frontend/src/pages/admin/BuildingManagement.js",
     "Create a reservation; within 30 s reload Reports + Building Management — both should reflect the new booking."),

    ("PPA-30", "User does not receive email notifications for waitlist / reservation events", "Notifications", "Medium", "Failed",
     "Backlog (P1)",
     "Email notifications for waitlisted users + reservation events is an explicit P1 future task in PRD.md. Backend currently writes in-app notifications. Bringing this to GA needs an SMTP/Resend/SendGrid integration (recommend Resend or Emergent SMTP). Not in scope for this defect cycle.",
     "Future: /app/backend/services/email.py (TBD)",
     "Out of scope for this validation cycle. Item is tracked in PRD.md backlog."),

    ("PPA-31", "Save button hidden under 'Made with Emergent' badge", "Admin Settings", "Medium", "Failed",
     "Resolved",
     "AdminLayout main content area now has pb-24/pb-28 (96px / 112px bottom padding). The platform's floating Emergent badge always sits in the bottom-right; this guarantees Save / submit buttons don't get covered.",
     "/app/frontend/src/pages/admin/AdminLayout.js (line 154-157)",
     "Open AdminSettings on a viewport ≥ md and verify the Save button is fully clickable; the badge no longer overlaps."),

    ("PPA-54", "Reports - date picker only allows changing END date", "Reports & Analytics", "Medium", "Failed",
     "Already Resolved",
     "Closed in QAT round (TCID-ANALYTICS-004). Calendar's onSelect now commits any non-null change immediately — single click works as a same-day range; second click extends or restarts the range.",
     "/app/frontend/src/pages/admin/Reports.js (Calendar onSelect)",
     "Open Reports → date picker → click any single date → it should be accepted as both start and end. Click another date → range expands."),

    ("PPA-57", "Recent Reservations table shows stale data", "Reports & Analytics", "Medium", "Failed",
     "Already Resolved",
     "Closed in QAT round (TCID-ANALYTICS-016). Reports.js now sorts client-side by created_at DESC and refetches on focus / visibility change. Recent Reservations always shows the latest activity.",
     "/app/frontend/src/pages/admin/Reports.js (fetchReservations + useEffect)",
     "Create a reservation; switch to Reports tab → it should appear at the top of Recent Reservations within 1 second of focus."),

    ("PPA-58", "User Management - tooltip overlaps with header", "User Management", "Low", "Failed",
     "Resolved",
     "Same universal fix as PPA-17. shadcn TooltipContent now uses Radix collisionPadding=16 and avoidCollisions=true.",
     "/app/frontend/src/components/ui/tooltip.jsx",
     "Hover the User Management header tooltip — it should flip downward instead of clashing with the page title."),

    ("PPA-59", "Block Slots dialog - only first conflicting slot reported", "Event Blocking", "Medium", "Failed",
     "Already Resolved",
     "Closed in QAT round (TCID-EVENT-BLOCKING-012). create_event_block now does a two-pass create: collect ALL conflicting slots in pass 1, raise once with the full comma-separated list, only insert in pass 2 if pass 1 was empty.",
     "/app/backend/routes/event_blocks.py (create_event_block)",
     "POST /api/event-blocks with 3 slot_ids that all conflict on the same date — error detail must list all 3 slot labels."),

    ("PPA-62", "Booking - able to cancel past reservations", "Reservations", "Medium", "Failed",
     "Resolved",
     "PUT /api/reservations/{id}/cancel now refuses if reservation date is strictly before today; returns 400 'Cannot cancel a past reservation'. Today's reservations remain cancelable. Verified via curl.",
     "/app/backend/routes/reservations.py (cancel_reservation, lines ~292-307)",
     "Pick a reservation with date < today, status in {pending, confirmed} → PUT /cancel returns 400."),

    ("PPA-64", "Building Management - slot label / status overlap", "Building Management", "Low", "Failed",
     "Verify (UI)",
     "Slot label rendering already uses `truncate` + `max-w-full` and a `title` attribute — so overflow gets ellipsis with full label on hover. If QA still sees overlap on a specific viewport, share screen size and we'll add a stricter min-width on the badge cluster.",
     "/app/frontend/src/pages/admin/BuildingManagement.js (lines 500-503)",
     "Test on viewports ≥ 1280px and < 768px. Report any specific overlap with a screenshot."),

    ("PPA-65", "Building Management - 'Active' character-limit indicator behavior", "Building Management", "Low", "Failed",
     "Already Resolved",
     "Closed in earlier QAT cycle. Phase 6 added explicit maxLength on building text inputs; PUT now uses a permissive BuildingUpdate so legacy buildings round-trip without 422 (TCID-BUILDING-MANAGEMENT-010).",
     "/app/backend/models/building.py, /app/frontend/src/pages/admin/BuildingManagement.js",
     "Try editing a legacy building whose name exceeds 30 chars — PUT should succeed; new buildings still constrained."),

    ("PPA-66", "Waitlist - Join Waitlist UX missing from booking flow", "Waitlist", "Medium", "Failed",
     "Verify (UI)",
     "Backend /api/waitlist endpoints all exist and were tested 14/14 in earlier cycles. The Join Waitlist button is rendered conditionally in BookingPage when slots are full and the building has waitlist_enabled=true. If QA cannot see it, confirm waitlist_enabled is true on the building's parking_config.",
     "/app/backend/routes/waitlist.py, /app/frontend/src/pages/BookingPage.js",
     "1) Set parking_configs.waitlist_enabled=true for the building; 2) Fully book all slots for a date; 3) Try to book → 'Join Waitlist' option should appear."),

    ("PPA-14b", "Reset User's Main Building - 'Building not found' error", "User Management", "Medium", "Failed",
     "Cannot Reproduce",
     "Could not reproduce on current build — PUT /api/users/{id} accepts main_building=null and assigned_buildings=[] cleanly in our smoke test. Likely an artifact of stale frontend state OR the user's current main_building no longer exists in DB. Recommend QA re-test on freshly deployed build and capture the exact request payload + response.",
     "/app/backend/routes/users.py (update_user)",
     "Re-run with browser DevTools network tab open and capture the failing PUT request body + 4xx response."),

    ("PPA-14c", "Block a slot from Building Management - status overlaps with name", "Building Management", "Low", "Failed",
     "Verify (UI)",
     "Same as PPA-64 — slot rendering uses truncate. Need a screenshot at a specific viewport size to reproduce.",
     "/app/frontend/src/pages/admin/BuildingManagement.js",
     "Reproduce with screenshot + viewport width."),

    ("PPA-14d", "View Admin's Dashboard (duplicate row, no detail)", "Admin Dashboard", "—", "Failed",
     "Cannot Reproduce",
     "Row in source spreadsheet had no Title or Actual Result populated — appears to be a duplicate / data-entry artifact. Need QA to specify what was being validated.",
     "—",
     "Please clarify the test scenario; row is currently empty."),

    ("PPA-14e", "View Admin's Dashboard (duplicate row, no detail)", "Admin Dashboard", "—", "Failed",
     "Cannot Reproduce",
     "Same as above — duplicate empty row.",
     "—",
     "Please clarify the test scenario; row is currently empty."),
]


# ---------- Build workbook ----------
wb = Workbook()
ws = wb.active
ws.title = "Defect Response"

HEADERS = ["Defect ID", "Title", "Module", "Severity", "Customer Status",
           "Resolution Status", "Remarks / Resolution Detail",
           "Code Path / Files", "QA Verification Steps"]

hdr_fill = PatternFill("solid", fgColor="08263E")
hdr_font = Font(color="FFFFFF", bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
thin = Side(border_style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = center; cell.border = border

for r, row in enumerate(ITEMS, 2):
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.alignment = left_wrap if c >= 7 else center
        cell.border = border
    sev = row[3]
    sc = ws.cell(row=r, column=4)
    sc.fill = PatternFill("solid", fgColor=SEV_COLORS.get(sev, "DDDDDD"))
    sc.font = Font(color="FFFFFF", bold=True); sc.alignment = center
    rstat = row[5]
    rc = ws.cell(row=r, column=6)
    rc.fill = PatternFill("solid", fgColor=STATUS_COLORS.get(rstat, "DDDDDD"))
    rc.font = Font(color="FFFFFF", bold=True); rc.alignment = center

WIDTHS = [10, 50, 22, 11, 18, 18, 65, 50, 50]
for i, w in enumerate(WIDTHS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 32
for r in range(2, len(ITEMS) + 2):
    ws.row_dimensions[r].height = 130
ws.freeze_panes = "A2"

# ---- Summary tab ----
sm = wb.create_sheet("Summary")
sm["A1"] = "Defect Validation — Response Summary"
sm["A1"].font = Font(size=16, bold=True, color="08263E")
sm["A2"] = f"Source: 'Parking App - Defect Validation.xlsx'   |   Generated {datetime.now().strftime('%B %d, %Y')}"
sm["A2"].font = Font(italic=True, color="666666")

sm["A4"] = "Total customer-reported defect rows"
sm["B4"] = len(ITEMS)
sm["A4"].font = Font(bold=True)

sm["A6"] = "Resolution status breakdown"
sm["A6"].font = Font(bold=True)
i = 7
for st in ["Resolved", "Already Resolved", "Verify Deploy", "Verify (UI)", "Backlog (P1)", "Cannot Reproduce", "Open"]:
    cell = sm.cell(row=i, column=1, value=st)
    cell.fill = PatternFill("solid", fgColor=STATUS_COLORS[st])
    cell.font = Font(color="FFFFFF", bold=True)
    sm.cell(row=i, column=2, value=sum(1 for x in ITEMS if x[5] == st))
    i += 1

i += 1
sm.cell(row=i, column=1, value="Items genuinely resolved this cycle (new code)").font = Font(bold=True, color="27AE60")
i += 1
for x in ITEMS:
    if x[5] == "Resolved":
        sm.cell(row=i, column=1, value=f"{x[0]} — {x[1]}").font = Font(bold=True)
        sm.cell(row=i, column=2, value=x[3])
        i += 1

sm.column_dimensions["A"].width = 75
sm.column_dimensions["B"].width = 30

wb.save(OUT)
print(f"Wrote {OUT}")
print(f"Total rows: {len(ITEMS)}")
print(f"  Resolved (this cycle): {sum(1 for x in ITEMS if x[5]=='Resolved')}")
print(f"  Already resolved (prior): {sum(1 for x in ITEMS if x[5]=='Already Resolved')}")
print(f"  Verify deploy/UI: {sum(1 for x in ITEMS if x[5] in ('Verify Deploy','Verify (UI)'))}")
print(f"  Backlog: {sum(1 for x in ITEMS if x[5]=='Backlog (P1)')}")
print(f"  Cannot reproduce: {sum(1 for x in ITEMS if x[5]=='Cannot Reproduce')}")
